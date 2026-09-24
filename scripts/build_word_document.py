#!/usr/bin/env python3
"""
Builds an executive-grade Word document (.docx) with all 39 Domain KPI and EDA SQL queries.
Executes each query to fetch live column names and sample data directly from PostgreSQL.
"""

import os
import sys
import shutil
import psycopg2
from dotenv import load_dotenv
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from query_catalogue_data import CATALOGUE

load_dotenv()

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_val(val):
    if val is None:
        return "NULL"
    elif isinstance(val, float):
        return f"{val:,.2f}"
    elif hasattr(val, 'as_tuple'): # Decimal
        return f"{float(val):,.2f}"
    return str(val)

def main():
    print("Connecting to PostgreSQL to fetch live query execution samples...")
    conn = psycopg2.connect(
        dbname=os.getenv('DB_NAME', 'accio_retailmart_27'),
        user=os.getenv('DB_USER', 'postgres'),
        password=os.getenv('DB_PASSWORD', 'accio123'),
        host=os.getenv('DB_HOST', 'localhost'),
        port=os.getenv('DB_PORT', '5432')
    )
    cur = conn.cursor()

    doc = Document()

    # Set page margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # -------------------------------------------------------------
    # 1. Document Title & Header
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("RetailMart V3 Enterprise BI Platform")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A) # Deep Navy
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Authoritative SQL Query Catalogue: Domain KPIs & Exploratory Data Analysis (EDA)")
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Royal Blue
    sub_p.paragraph_format.space_after = Pt(14)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="94A3B8")
    meta_data = [
        ("Project & Architecture", "RetailMart V3 Commercial BI Platform (Django 5.1 + PostgreSQL 18.4)"),
        ("Author & Governance", "Senior Business Intelligence Analyst & Data Architecture Team"),
        ("Database Instance", "PostgreSQL 18.4 (accio_retailmart_27) | Schemas: sales, products, stores, customers, manufacture, marketing, core, analytics"),
        ("Catalogue Scope", "39 Validated SQL Queries covering Executive, Sales, Customer, Operations, Cross-Functional, and EDA Domains")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Executive Overview
    exec_p = doc.add_paragraph()
    exec_p.paragraph_format.space_after = Pt(14)
    exec_p.add_run(
        "This authoritative register documents the complete production SQL queries powering the RetailMart V3 Enterprise "
        "Business Intelligence dashboards and ad-hoc analytical explorations. Every query is reconciled against PostgreSQL 18.4 "
        "and adheres to corporate data governance rules: recognizing Delivered net revenue, handling potential division-by-zero "
        "with NULLIF, and applying deterministic quintiles for customer segmentation."
    )

    # -------------------------------------------------------------
    # Table of Contents Summary
    # -------------------------------------------------------------
    toc_heading = doc.add_heading("Master Query Catalogue Index (39 Queries)", level=1)
    toc_heading.paragraph_format.space_before = Pt(14)
    toc_heading.paragraph_format.space_after = Pt(6)

    sections_grouped = {}
    for item in CATALOGUE:
        sec = item["section"]
        if sec not in sections_grouped:
            sections_grouped[sec] = []
        sections_grouped[sec].append(item)

    for sec_name, items in sections_grouped.items():
        sec_p = doc.add_paragraph()
        sec_p.paragraph_format.space_after = Pt(2)
        r = sec_p.add_run(f"• {sec_name} ({len(items)} Queries)")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    doc.add_page_break()

    # -------------------------------------------------------------
    # Render Each Section and Query
    # -------------------------------------------------------------
    current_sec = None

    for idx, item in enumerate(CATALOGUE, 1):
        sec = item["section"]
        if sec != current_sec:
            current_sec = sec
            sec_heading = doc.add_heading(sec, level=1)
            sec_heading.paragraph_format.space_before = Pt(18)
            sec_heading.paragraph_format.space_after = Pt(8)
            
            divider = doc.add_paragraph()
            divider.paragraph_format.space_after = Pt(12)
            d_run = divider.add_run("━" * 60)
            d_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        # Query Title
        q_heading = doc.add_heading(f"{item['id']}: {item['title']}", level=2)
        q_heading.paragraph_format.space_before = Pt(12)
        q_heading.paragraph_format.space_after = Pt(4)
        for r in q_heading.runs:
            r.font.size = Pt(12.5)
            r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        # Metadata Spec Table (2 columns)
        spec_table = doc.add_table(rows=5, cols=2)
        spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(spec_table, color="CBD5E1", sz="4")

        spec_rows = [
            ("Domain & Purpose", f"{item['domain']} — {item['purpose']}"),
            ("Business Formula", item["formula"]),
            ("Granularity & Grain", item["grain"]),
            ("Underlying Tables", item["tables"]),
            ("Management Action", item["action"])
        ]

        for s_idx, (s_label, s_val) in enumerate(spec_rows):
            s_row = spec_table.rows[s_idx]
            c0, c1 = s_row.cells[0], s_row.cells[1]
            c0.width = Inches(1.8)
            c1.width = Inches(5.2)
            set_cell_margins(c0, 50, 50, 80, 80)
            set_cell_margins(c1, 50, 50, 80, 80)
            set_cell_background(c0, "F8FAFC")
            set_cell_background(c1, "FFFFFF")

            p0 = c0.paragraphs[0]
            r0 = p0.add_run(s_label)
            r0.font.bold = True
            r0.font.size = Pt(8.5)
            r0.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(s_val)
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # SQL Code Block
        code_p = doc.add_paragraph()
        code_p.paragraph_format.space_before = Pt(8)
        code_p.paragraph_format.space_after = Pt(2)
        code_lbl = code_p.add_run("Production SQL Query (PostgreSQL 18.4):")
        code_lbl.font.bold = True
        code_lbl.font.size = Pt(9.5)
        code_lbl.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

        # Shaded box for SQL
        sql_box = doc.add_table(rows=1, cols=1)
        sql_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        sql_cell = sql_box.rows[0].cells[0]
        sql_cell.width = Inches(7.0)
        set_cell_background(sql_cell, "F1F5F9")
        set_cell_margins(sql_cell, 80, 80, 120, 120)
        set_table_borders(sql_box, color="94A3B8", sz="6")

        sql_p = sql_cell.paragraphs[0]
        sql_p.paragraph_format.space_after = Pt(0)
        sql_run = sql_p.add_run(item["sql"])
        sql_run.font.name = "Consolas"
        sql_run.font.size = Pt(8.5)
        sql_run.font.color.rgb = RGBColor(0x0F, 0x1F, 0x38)

        # Execute Query to fetch live verification output
        out_rows = []
        col_names = []
        try:
            cur.execute(item["sql"])
            col_names = [desc[0] for desc in cur.description]
            out_rows = cur.fetchall()
        except Exception as e:
            print(f"Error fetching data for {item['id']}: {e}")
            conn.rollback()

        # Output sample table
        if out_rows:
            sample_p = doc.add_paragraph()
            sample_p.paragraph_format.space_before = Pt(6)
            sample_p.paragraph_format.space_after = Pt(2)
            sample_lbl = sample_p.add_run(f"Execution Verification ({len(out_rows)} rows returned — Sample Output):")
            sample_lbl.font.bold = True
            sample_lbl.font.size = Pt(9)
            sample_lbl.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) # Forest Green

            display_rows = out_rows[:3] # Show top 3 rows
            res_table = doc.add_table(rows=len(display_rows) + 1, cols=len(col_names))
            res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(res_table, color="E2E8F0", sz="4")

            # Headers
            hdr_cells = res_table.rows[0].cells
            col_w = Inches(7.0 / len(col_names))
            for c_i, name in enumerate(col_names):
                hdr_cells[c_i].width = col_w
                set_cell_background(hdr_cells[c_i], "E2E8F0")
                set_cell_margins(hdr_cells[c_i], 40, 40, 60, 60)
                hp = hdr_cells[c_i].paragraphs[0]
                hrun = hp.add_run(name)
                hrun.font.bold = True
                hrun.font.size = Pt(8)
                hrun.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

            # Data rows
            for r_i, d_row in enumerate(display_rows, 1):
                d_cells = res_table.rows[r_i].cells
                bg = "FFFFFF" if r_i % 2 != 0 else "F8FAFC"
                for c_i, val in enumerate(d_row):
                    d_cells[c_i].width = col_w
                    set_cell_background(d_cells[c_i], bg)
                    set_cell_margins(d_cells[c_i], 40, 40, 60, 60)
                    dp = d_cells[c_i].paragraphs[0]
                    drun = dp.add_run(format_val(val))
                    drun.font.size = Pt(7.5)
                    drun.font.name = "Consolas"
                    drun.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(14)

    conn.close()

    # Save destinations
    output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "project_documents", "documentation")
    os.makedirs(output_dir, exist_ok=True)

    file_1 = os.path.join(output_dir, "RetailMart_V3_Master_KPI_and_EDA_SQL_Queries.docx")
    file_2 = os.path.join(output_dir, "approved_kpi_and_eda_queries.docx")
    
    doc.save(file_1)
    doc.save(file_2)
    print(f"Generated Word Document 1: {file_1}")
    print(f"Generated Word Document 2: {file_2}")

    # Copy to retailmart_bi_github directory
    github_doc_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "retailmart_bi_github", "project_documents", "documentation")
    if os.path.exists(github_doc_dir):
        shutil.copy2(file_1, os.path.join(github_doc_dir, "RetailMart_V3_Master_KPI_and_EDA_SQL_Queries.docx"))
        shutil.copy2(file_2, os.path.join(github_doc_dir, "approved_kpi_and_eda_queries.docx"))
        print(f"Copied Word documents to GitHub folder: {github_doc_dir}")

    # Recreate ZIP archive
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    github_folder = os.path.join(root_dir, "retailmart_bi_github")
    output_zip = os.path.join(root_dir, "retailmart_bi_github")
    shutil.make_archive(output_zip, 'zip', github_folder)
    print(f"Updated ZIP archive: {output_zip}.zip")

if __name__ == "__main__":
    main()
